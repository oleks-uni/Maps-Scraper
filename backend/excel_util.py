from io import BytesIO
import openpyxl
from openpyxl.utils import get_column_letter

def results_to_xlsx_bytes(results: list[dict]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Places"

    headers = ["name", "rating_info", "address"]
    ws.append(headers)

    for item in results:
        ws.append([
            item.get("name"),
            item.get("rating_info"),
            item.get("address"),
        ])

    # Ширина колонок + заморозити заголовок
    ws.freeze_panes = "A2"
    for i, _ in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = 35

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()